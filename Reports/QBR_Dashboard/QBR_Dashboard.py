
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
import re

# Project root: Coreweave Outstanding Reports
PROJECT_FOLDER = Path(__file__).resolve().parents[2]

# Dashboard paths
DASHBOARD_FOLDER = PROJECT_FOLDER / "Reports" / "QBR_Dashboard"
INPUT_FOLDER = DASHBOARD_FOLDER / "Input"
OUTPUT_FOLDER = DASHBOARD_FOLDER / "Output"

# Input and output files
ADDRESS_FILE = PROJECT_FOLDER / "Address Cache" / "address_input.xlsx"
OUTPUT_FILE = OUTPUT_FOLDER / "Cases.xlsx"

# Salesforce exports
REPORT_NAME_PATTERN = "CoreWeave QBR Dashboard-*.xlsx"

# Validation settings
INVALID_PART_TYPES = {"NUM", "ABC", "XYZ", "123", "456"}
VALID_GPU_MODELS = {"A100", "B200", "H100", "H200"}

EXPECTED_COLUMNS = [
    "Case Number",
    "Case Creation Date",
    "Resolution Date/Time",
    "Status",
    "Work Type",
    "Part Type",
    "Faulty Part Number",
    "GPU Model",
    "GPU Failure Codes",
    "Issue Sub-Type",
    "Issue Sub-Type Detail",
    "Resolution Category",
    "Resolution Code",
    "Asset Location: Location Name",
    "System SKU",
    "Root Cause: Root Cause Name",
    "City",
    "State",
    "Region",
]

# Required source fields
REQUIRED_COLUMNS = [
    "Case Number",
    "Case Creation Date",
    "Resolution Date/Time",
    "Status",
    "Work Type",
    "Faulty Part Number",
]

# Salesforce field mapped to the Power BI System SKU output column
SALESFORCE_SKU_COLUMN = "Asset: Product: Product Name"

INPUT_FOLDER.mkdir(parents=True, exist_ok=True)
OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)


def normalize_case_number(series):
    """Normalize Salesforce case IDs for reliable matching."""
    return series.astype("string").str.strip()


def extract_gpu_model(value):
    """Extract a supported GPU model from a faulty part number."""
    if pd.isna(value):
        return ""

    parts = str(value).upper().split("-")

    for model in VALID_GPU_MODELS:
        if model in parts:
            return model

    return ""




def clean_text(value):
    """Convert missing values to empty strings and trim whitespace."""
    if pd.isna(value):
        return ""
    return str(value).strip()


def parse_asset_address(value):
    """
    Best-effort fallback parser for Salesforce asset addresses.

    Returns (City, State, Region).
    Uncertain fields are left blank rather than guessed.
    """
    address = clean_text(value)

    if not address:
        return "", "", ""

    # Normalize whitespace.
    address = re.sub(r"\s+", " ", address).strip()
    upper = address.upper()


    # US addresses: ... Phoenix, AZ 85027-5139, USA
    us_match = re.search(
        r",\s*([^,]+),\s*([A-Z]{2})\s+\d{5}(?:-\d{4})?"
        r"(?:,\s*USA)?$",
        upper,
    )

    if us_match:
        city = us_match.group(1).strip().title()
        state = us_match.group(2).strip()

        # Preserve the existing regional convention.
        west_coast = {"CA", "OR", "WA"}

        # Match the existing Address Cache regional conventions.
        if state in {"CA", "OR"}:
            region = "West Coast (US)"
        elif state in {"WA"}:
            region = "West Coast (US)"
        elif state in {"TX", "IL", "IN", "AZ"}:
            region = "Central (US)"
        elif state in {"MA", "GA", "OH"}:
            region = "East Coast (US)"
        else:
            region = ""

        return city, state, region


     # Sweden: ... 791 77 Falun, Sweden
    sweden_match = re.search(
        r"\b\d{3}\s?\d{2}\s+([^,]+),\s*SWEDEN$",
        upper,
    )
    if sweden_match:
        return sweden_match.group(1).strip().title(), "Sweden", "Europe"

    # UK: ... London E14 2AX, UK
    uk_match = re.search(
        r",\s*([^,]+?)\s+[A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2},\s*UK$",
        upper,
    )
    if uk_match:
        return uk_match.group(1).strip().title(), "UK", "Europe"

    # Spain: supports both "08840 Viladecans, Barcelona, Spain"
    # and "08040 Barcelona, Spain".
    spain_match = re.search(
        r"\b(\d{5})\s+([^,]+)(?:,\s*[^,]+)?,\s*SPAIN$",
        upper,
    )

    if spain_match:
        city = spain_match.group(2).strip().title()
        return city, "Spain", "Europe"

    # Unknown or ambiguous format: do not invent values.
    return "", "", ""

def process_cases(input_file, address_file):
    # Salesforce report has 10 rows before the column headers.
    cases = pd.read_excel(input_file, skiprows=10)
    cases = cases.dropna(axis=1, how="all")
    cases.columns = cases.columns.astype(str).str.strip()

    # Ensure the optional asset-address column exists for fallback parsing.
    asset_address_column = "Asset Location: Location Name"

    if asset_address_column not in cases.columns:
        cases[asset_address_column] = ""

    # Check for duplicate source headers.
    duplicate_columns = cases.columns[
        cases.columns.duplicated()
    ].tolist()

    if duplicate_columns:
        raise ValueError(
            f"Duplicate Salesforce column names: {duplicate_columns}"
        )

    missing = [
        column for column in REQUIRED_COLUMNS
        if column not in cases.columns
    ]

    if missing:
        raise ValueError(
            f"Missing required Salesforce columns: {missing}\n"
            f"Actual columns found: {cases.columns.tolist()}"
        )

    # Normalize case IDs before joining.
    cases["Case Number"] = normalize_case_number(cases["Case Number"])

    # Map Salesforce's asset product name to System SKU.
    if SALESFORCE_SKU_COLUMN in cases.columns:
        cases["System SKU"] = cases[SALESFORCE_SKU_COLUMN]
        print(
            f"\nSystem SKU mapped from: {SALESFORCE_SKU_COLUMN}"
        )
    elif "System SKU" not in cases.columns:
        cases["System SKU"] = ""
        print(
            "\nWARNING: Neither the Salesforce asset product field "
            "nor System SKU was found."
        )

    # Convert dates while supporting mixed source formats.
    for column in ["Case Creation Date", "Resolution Date/Time"]:
        cases[column] = pd.to_datetime(
            cases[column],
            format="mixed",
            errors="coerce",
        )

    # Derive Part Type from the first three characters of the part number.
    faulty_parts = cases["Faulty Part Number"].fillna("").astype(str)
    cases["Part Type"] = faulty_parts.str[:3].str.upper()

    cases.loc[
        cases["Part Type"].isin(INVALID_PART_TYPES),
        "Part Type",
    ] = ""

    # Derive GPU Model without changing the original part number.
    cases["GPU Model"] = cases["Faulty Part Number"].apply(
        extract_gpu_model
    )

    print("\nGPU Model counts:")
    print(cases["GPU Model"].value_counts(dropna=False).to_string())

    print("\nPart Type counts:")
    print(cases["Part Type"].value_counts(dropna=False).to_string())

    # Load address cache.
    addresses = pd.read_excel(address_file)
    addresses.columns = addresses.columns.astype(str).str.strip()

    address_columns = ["Case Number", "City", "State", "Region"]
    missing = [
        column for column in address_columns
        if column not in addresses.columns
    ]

    if missing:
        raise ValueError(
            f"Missing columns in address cache: {missing}"
        )

    addresses = addresses[address_columns].copy()
    addresses["Case Number"] = normalize_case_number(
        addresses["Case Number"]
    )

    # Remove blank case IDs from the address cache.
    addresses = addresses[
        addresses["Case Number"].notna()
        & addresses["Case Number"].ne("")
    ].copy()

    # Prevent duplicate cache entries from multiplying case rows.
    duplicates = addresses[
        addresses["Case Number"].duplicated(keep=False)
    ]

    if not duplicates.empty:
        raise ValueError(
            "Address cache contains duplicate Case Numbers:\n"
            f"{duplicates['Case Number'].unique()[:20].tolist()}"
        )

    # Address cache supplies City, State, and Region.
    cases = cases.drop(
        columns=["City", "State", "Region"],
        errors="ignore",
    )

    original_count = len(cases)

    cases = cases.merge(
        addresses,
        on="Case Number",
        how="left",
        validate="many_to_one",
        indicator=True,
        sort=False,
    )

    if len(cases) != original_count:
        raise RuntimeError(
            "Case count changed during the address merge."
        )


    # Fill missing address fields using the Salesforce asset address.
    asset_address_column = "Asset Location: Location Name"

    unmatched_mask = cases["_merge"].eq("left_only")
    unmatched_count = int(unmatched_mask.sum())
    fallback_count = 0

    print(f"\nCases processed: {len(cases)}")
    print(f"Cases with no address-cache match: {unmatched_count}")

    for index in cases.index:
        asset_address = cases.at[index, asset_address_column] \
            if asset_address_column in cases.columns else ""

        parsed_city, parsed_state, parsed_region = parse_asset_address(
            asset_address
        )

        for column, parsed_value in [
            ("City", parsed_city),
            ("State", parsed_state),
            ("Region", parsed_region),
        ]:
            existing_value = clean_text(cases.at[index, column])

            # Preserve existing cache values; only fill blanks.
            if not existing_value and parsed_value:
                cases.at[index, column] = parsed_value
                fallback_count += 1

    print(f"Address fields filled by fallback: {fallback_count}")

    # Save cases that still have any missing address field.
    missing_address_mask = (
        cases[["City", "State", "Region"]]
        .fillna("")
        .astype(str)
        .apply(lambda column: column.str.strip().eq(""))
        .any(axis=1)
    )

    missing_addresses = cases.loc[
        missing_address_mask,
        [
            "Case Number",
            asset_address_column,
            "City",
            "State",
            "Region",
        ],
    ].copy()

    if not missing_addresses.empty:
        unmatched_file = OUTPUT_FOLDER / "Cases_Missing_Addresses.xlsx"
        missing_addresses.to_excel(unmatched_file, index=False)
        print(
            f"Cases still missing at least one address field: "
            f"{len(missing_addresses)}"
        )
        print(f"Review file saved: {unmatched_file}")

    cases = cases.drop(columns=["_merge"])


    # Report optional columns absent from the Salesforce export.
    # System SKU has already been mapped above if its source exists.
    optional_source_columns = [
        "Issue Sub-Type Detail",
        "Resolution Category",
        "Resolution Code",
        "Asset Location: Location Name",
        "Root Cause: Root Cause Name",
    ]

    missing_optional = [
        column for column in optional_source_columns
        if column not in cases.columns
    ]

    if missing_optional:
        print(
            "\nWARNING: These source columns are missing and "
            "will be blank:"
        )
        for column in missing_optional:
            print(f"  - {column}")

    # Preserve the expected Power BI schema.
    for column in EXPECTED_COLUMNS:
        if column not in cases.columns:
            cases[column] = ""

    cases = cases[EXPECTED_COLUMNS]

    # Export with the worksheet name expected by Power BI.
    try:
        cases.to_excel(
            OUTPUT_FILE,
            index=False,
            sheet_name="Cases",
        )
    except PermissionError:
        raise PermissionError(
            f"Cannot write to {OUTPUT_FILE}. "
            "Close Cases.xlsx in Excel or Power BI and retry."
        ) from None

    # Apply Excel date display formatting.
    workbook = load_workbook(OUTPUT_FILE)
    sheet = workbook["Cases"]

    for column_name in ["Case Creation Date", "Resolution Date/Time"]:
        column_number = cases.columns.get_loc(column_name) + 1

        for row in range(2, sheet.max_row + 1):
            sheet.cell(
                row=row,
                column=column_number,
            ).number_format = "mm/dd/yy"

    workbook.save(OUTPUT_FILE)

    print(f"\nSaved: {OUTPUT_FILE}")

    for column in ["City", "State", "Region"]:
        missing_count = (
            cases[column]
            .fillna("")
            .astype(str)
            .str.strip()
            .eq("")
            .sum()
        )
        print(f"Cases missing {column}: {missing_count}")

    for column in [
        "Issue Sub-Type Detail",
        "Resolution Category",
        "Resolution Code",
        "Asset Location: Location Name",
        "System SKU",
        "Root Cause: Root Cause Name",
    ]:
        populated = (
            cases[column]
            .fillna("")
            .astype(str)
            .str.strip()
            .ne("")
            .sum()
        )
        print(f"{column} populated: {populated}")

    return OUTPUT_FILE


if __name__ == "__main__":
    files = list(INPUT_FOLDER.glob(REPORT_NAME_PATTERN))

    if not files:
        raise FileNotFoundError(
            f"No matching QBR export found in {INPUT_FOLDER}"
        )

    # Process the most recently modified Salesforce export.
    input_file = max(files, key=lambda file: file.stat().st_mtime)

    process_cases(
        input_file=input_file,
        address_file=ADDRESS_FILE,
    )
